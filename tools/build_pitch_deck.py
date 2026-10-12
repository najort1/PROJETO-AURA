"""Gera o pitch deck executivo de 5 min do Projeto Aura 2.0 (Azul Fidelidade).

Editável por design: python-pptx com textos, tabelas e DOIS gráficos nativos
(Pareto de concentração e ΔReceita por cenário). Tema institucional Azul:
tokens de azul-design-patterns.md. Reproduzível: `python tools/build_pitch_deck.py`.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "pitch_deck_aura2.pptx"

# Tokens Azul (azul-design-patterns.md)
AZUL_ACTION = RGBColor(0x08, 0x74, 0xB5)
AZUL_PRIMARY = RGBColor(0x00, 0x8B, 0xC4)
AZUL_DEEP = RGBColor(0x03, 0x37, 0x74)
AZUL_NAVY = RGBColor(0x0B, 0x35, 0x58)
MAGENTA = RGBColor(0xE8, 0x32, 0xA5)
LIME = RGBColor(0xD8, 0xED, 0x00)
CANVAS = RGBColor(0xF8, 0xF9, 0xFA)
BORDER = RGBColor(0xE2, 0xE6, 0xE9)
TEXT_MAIN = RGBColor(0x1E, 0x2A, 0x30)
TEXT_SOFT = RGBColor(0x65, 0x76, 0x81)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Inches(13.333), Inches(7.5)


def _slide(prs: Presentation) -> object:
    return prs.slides.add_slide(prs.slide_layouts[6])


def _bg(slide: object, color: RGBColor) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _box(slide: object, x: float, y: float, w: float, h: float) -> object:
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def _para(tf: object, text: str, size: int, color: RGBColor, bold: bool = False, align: object = PP_ALIGN.LEFT) -> None:
    p = tf.paragraphs[0] if not tf.paragraphs[0].runs and not tf.paragraphs[0].text else tf.add_paragraph()
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Calibri"
    p.alignment = align


def _kpi(slide: object, x: float, y: float, value: str, label: str, accent: RGBColor = AZUL_ACTION) -> None:
    tf = _box(slide, x, y, 2.9, 1.5)
    _para(tf, value, 26, accent, bold=True)
    _para(tf, label, 11, TEXT_SOFT)


def _stripe(slide: object, y: float = 0.0, color: RGBColor = MAGENTA, h: float = 0.06) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(y), W, Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def _footer(slide: object, note: str) -> None:
    tf = _box(slide, 0.6, 6.95, 12.1, 0.4)
    _para(tf, note, 9, TEXT_SOFT)


def build() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # ---------------- 1 · CAPA ----------------
    s = _slide(prs)
    _bg(s, AZUL_DEEP)
    _stripe(s, y=0.55, color=MAGENTA)
    tf = _box(s, 0.9, 2.1, 11.5, 1.2)
    _para(tf, "Projeto Aura 2.0", 54, WHITE, bold=True)
    tf = _box(s, 0.9, 3.35, 11.5, 1.6)
    _para(tf, "Da milha ao momento de cuidado", 28, AZUL_PRIMARY)
    _para(tf, "Inteligência de clientes e retenção — Azul Fidelidade", 16, RGBColor(0x9A, 0xE0, 0xE8))
    tf = _box(s, 0.9, 6.1, 11.5, 1.0)
    _para(
        tf,
        "Pitch executivo de 5 minutos · Pipeline 100% reproduzível por código · Base sintética calibrada com fontes públicas ANAC",
        12,
        RGBColor(0x9A, 0xE0, 0xE8),
    )
    _footer(s, "Dados sintéticos — nenhum dado pessoal real. Métricas validam o pipeline; não são prova de desempenho no mundo real.")

    # ---------------- 2 · PARADOXO ----------------
    s = _slide(prs)
    _bg(s, CANVAS)
    _stripe(s, color=AZUL_ACTION)
    tf = _box(s, 0.6, 0.35, 12.1, 0.9)
    _para(tf, "1 · O paradoxo da lealdade", 34, AZUL_DEEP, bold=True)
    tf = _box(s, 0.6, 1.3, 6.0, 4.4)
    _para(tf, "20,1 milhões de membros no Azul Fidelidade (RI, dez/2025).", 18, TEXT_MAIN)
    _para(tf, "Mas a carteira real de fidelidade é concentrada e dormente:", 18, TEXT_MAIN)
    _para(tf, "• 48% dos clientes são promocionais ocasionais (churn previsto 74%)", 15, TEXT_MAIN)
    _para(tf, "• 10% dos clientes carregam 47% da receita prevista de 12 meses", 15, TEXT_MAIN)
    _para(tf, "• R$ 8,31 milhões em valor em risco na carteira simulada (CLV × P(churn))", 15, TEXT_MAIN)
    _para(tf, "Manter quem importa custa uma fração de conquistar quem não existe:", 15, TEXT_SOFT)
    _para(tf, "o problema não é falta de dados — é não saber QUEM priorizar, COM QUER cuidar e QUANDO agir.", 15, TEXT_SOFT)
    _kpi(s, 7.1, 1.6, "20,1M", "membros (fato RI dez/2025)")
    _kpi(s, 10.0, 1.6, "50.000", "clientes na base sintética calibrada v3")
    _kpi(s, 7.1, 3.3, "48,4%", "persona eventual promocional (churn 73,6%)", accent=MAGENTA)
    _kpi(s, 10.0, 3.3, "R$ 8,31M", "valor em risco simulado (Σ CLV × churn)")
    _footer(s, "Números da base sintética v3 (validação de pipeline); escala de membros verificada por pesquisa no RI da Azul.")

    # ---------------- 3 · COMO IDENTIFICAMOS VALOR ----------------
    s = _slide(prs)
    _bg(s, CANVAS)
    _stripe(s, color=AZUL_ACTION)
    tf = _box(s, 0.6, 0.35, 12.1, 0.9)
    _para(tf, "2 · Valor é CLV previsto — não gasto passado", 34, AZUL_DEEP, bold=True)
    tf = _box(s, 0.6, 1.25, 5.6, 5.0)
    _para(tf, "Pipeline (tudo auditável):", 16, TEXT_MAIN, bold=True)
    _para(tf, "• Fontes públicas ANAC (3,2 GB) → calibração com manifesto sha256", 13, TEXT_MAIN)
    _para(tf, "• Base sintética determinística: ground truth segregado, zero PII", 13, TEXT_MAIN)
    _para(tf, "• Features anti-vazamento temporal (cutoff 2025-08-31, testes)", 13, TEXT_MAIN)
    _para(tf, "• CLV: LightGBM WAPE 0,655 vs 0,797 do baseline BG/NBD+GG (Δ −0,142, IC 95%)", 13, TEXT_MAIN)
    _para(tf, "• Churn: ROC-AUC 0,857 · PR-AUC 0,852 · Brier 0,145", 13, TEXT_MAIN)
    _para(tf, "• ScoreAV (0,6 CLV · 0,2 RFM · 0,2 potencial): Spearman > 0,998 sob ±10% nos pesos", 13, TEXT_MAIN)
    _para(tf, "• Explicação SHAP por cliente — toda decisão tem 'por quê'", 13, TEXT_MAIN)
    # Gráfico nativo: Pareto
    data = CategoryChartData()
    data.categories = ["Top 5%", "Top 10%", "Top 15%", "Top 20%"]
    data.add_series("% da receita prevista capturada", (0.3049, 0.4696, 0.5739, 0.6488))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.6), Inches(1.5), Inches(6.1), Inches(4.6), data)
    chart = gf.chart
    chart.has_legend = False
    chart.chart_title.text_frame.text = "Concentração de receita — corte Top 10% captura 46,96%"
    _footer(s, "Corte alto-valor: 5.005 clientes (10,01% da carteira) capturando 46,96% da receita. Validade: métricas computadas na base sintética com validação temporal.")

    # ---------------- 4 · DEMO ----------------
    s = _slide(prs)
    _bg(s, CANVAS)
    _stripe(s, color=AZUL_ACTION)
    tf = _box(s, 0.6, 0.35, 12.1, 0.9)
    _para(tf, "3 · Do risco à mensagem em 30 ms (demo ao vivo)", 34, AZUL_DEEP, bold=True)
    tf = _box(s, 0.6, 1.25, 5.9, 5.2)
    _para(tf, "Cliente real da base: C038454", 17, TEXT_MAIN, bold=True)
    _para(tf, "• Topázio · persona Eventual Ascendente", 14, TEXT_MAIN)
    _para(tf, "• CLV previsto R$ 4.975 · churn 62,5% · VaR R$ 3.109", 14, TEXT_MAIN)
    _para(tf, "• SHAP: nível Topázio (+R$ 3.333) · ticket médio (+R$ 512) · pontos 12M (+R$ 346)", 14, TEXT_MAIN)
    _para(tf, "O agente compõe (offline, determinístico):", 17, TEXT_MAIN, bold=True)
    _para(tf, "• Ação: antecipação de nível com citação do regulamento (AF-TIER-TOPAZIO, snapshot 2026-10-12, status verificado)", 14, TEXT_MAIN)
    _para(tf, "• Canal, tom e benefício vêm do inventário elegível — nada é inventado", 14, TEXT_MAIN)
    _para(tf, "• Latência média 33 ms/cliente · custo R$ 0 (offline)", 14, TEXT_MAIN)
    _kpi(s, 7.0, 1.6, "62/62", "vetores adversariais bloqueados (4 rodadas)", accent=MAGENTA)
    _kpi(s, 10.0, 1.6, "100%", "groundedness das citações (170 clientes reais)")
    _kpi(s, 7.0, 3.3, "0", "violações de política/LGPD em 50.000 recomendações")
    _kpi(s, 10.0, 3.3, "44,7%", "human review no silêncio responsável (fail-safe)")
    _footer(s, "70,7% da carteira não recebe oferta (não há ação elegível honesta) — silêncio é decisão, não omissão. Modo LLM: plugável, auditável.")

    # ---------------- 5 · IMPACTO (HONESTO) ----------------
    s = _slide(prs)
    _bg(s, CANVAS)
    _stripe(s, color=MAGENTA)
    tf = _box(s, 0.6, 0.35, 12.1, 0.9)
    _para(tf, "4 · Impacto estimado — com a verdade na mesa", 34, AZUL_DEEP, bold=True)
    tf = _box(s, 0.6, 1.25, 5.9, 5.2)
    _para(tf, "O achado impopular (e valioso):", 16, TEXT_MAIN, bold=True)
    _para(tf, "Campanha generalista sobre 14.645 alvos é DEFICITÁRIA em todos os cenários —", 14, TEXT_MAIN)
    _para(tf, "o custo por ação (R$3–R$60) excede o retorno esperado em clientes de baixo valor em risco.", 14, TEXT_MAIN)
    _para(tf, "Break-even exige custo ≤ R$ 0,88/ação (base): só canais digitais de custo marginal ~0.", 14, TEXT_MAIN)
    _para(tf, "Estratégia que o número sustenta:", 16, TEXT_MAIN, bold=True)
    _para(tf, "• Priorização por VaR: R$ 80k cobrem 6.140 clientes e 58,9% do risco", 14, TEXT_MAIN)
    _para(tf, "• Piloto A/B (poder 80%, α 5%): n = 4.624/braço para detectar +2pp de retenção", 14, TEXT_MAIN)
    _para(tf, "• Na escala real (>140 mil alto-valor): 7+ pilotos simultâneos de +2pp", 14, TEXT_MAIN)
    # Gráfico nativo: ΔReceita por cenário
    data = CategoryChartData()
    data.categories = ["Conservador", "Base", "Otimista"]
    data.add_series("ΔReceita líquida (R$ mil)", (-163.4, -153.7, -132.9))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.6), Inches(1.5), Inches(6.1), Inches(4.6), data)
    chart = gf.chart
    chart.has_legend = False
    chart.chart_title.text_frame.text = "ΔReceita líquida por cenário — blanket é deficitário (IC95 ±R$2k)"
    _footer(s, "Efeito da ação (5–15%) e resposta (12–25%) são PREMISSAS de literatura CRM — não efeito causal. Uplift real exige o piloto.")

    # ---------------- 6 · RESPONSABILIDADE ----------------
    s = _slide(prs)
    _bg(s, CANVAS)
    _stripe(s, color=AZUL_ACTION)
    tf = _box(s, 0.6, 0.35, 12.1, 0.9)
    _para(tf, "5 · IA responsável como código, não como slide", 34, AZUL_DEEP, bold=True)
    tf = _box(s, 0.6, 1.25, 12.1, 5.2)
    _para(tf, "• Zero PII: base 100% sintética; para dados reais, RIPD e base legal antes de qualquer contato (LGPD art. 20: explicação + humano no circuito)", 15, TEXT_MAIN)
    _para(tf, "• Atributos sensíveis PROIBIDOS nas features — teste automatizado quebra o pipeline se vazar", 15, TEXT_MAIN)
    _para(tf, "• Auditoria Four-Fifths executada e PUBLICADA com 5 grupos fora do limiar (18-24: 0,795; Sul 0,558; Sudeste 0,642; C-Oeste 0,699; Nordeste 0,724) — medir, divulgar e mitigar, nunca maquiar", 15, TEXT_MAIN)
    _para(tf, "• 'Não informado' é rótulo honesto: ninguém ganha demografia inventada", 15, TEXT_MAIN)
    _para(tf, "• Guardrails: 69 testes permanentes de bloqueio (benefício fabricado, urgência manipulativa, dado sensível, injeção de prompt em PT/EN, SQL/XSS)", 15, TEXT_MAIN)
    _para(tf, "• Human-in-the-loop em toda oferta; registro auditável de decisões (decisions.jsonl)", 15, TEXT_MAIN)
    _para(tf, "• 190 testes automatizados verdes · 12 revisões independentes documentadas com veredito assinado", 15, TEXT_MAIN)
    _footer(s, "Relatórios: docs/responsible_ai.md · docs/ripd.md · auditoria reprodutível via tasks.py fairness")

    # ---------------- 7 · DADOS REAIS ----------------
    s = _slide(prs)
    _bg(s, CANVAS)
    _stripe(s, color=AZUL_ACTION)
    tf = _box(s, 0.6, 0.35, 12.1, 0.9)
    _para(tf, "6 · Plugar nos dados reais: troca de fonte, não reescrita", 34, AZUL_DEEP, bold=True)
    tf = _box(s, 0.6, 1.25, 5.9, 5.0)
    _para(tf, "• Contratos idênticos: 22 tabelas com schema validado em CI; dicionário gerado dos schemas", 14, TEXT_MAIN)
    _para(tf, "• Snowflake: DDL + COPY INTO prontos (sql/snowflake) · plano de transição em docs/snowflake_plan.md", 14, TEXT_MAIN)
    _para(tf, "• Ingestão versionada com manifesto sha256 — nova fonte entra com a mesma governança", 14, TEXT_MAIN)
    _para(tf, "• Agente offline (R$ 0) é o padrão auditado; LLM pluggável (OpenAI-compatible) custa R$ 66–1.238 por 50 mil recomendações conforme o tier", 14, TEXT_MAIN)
    _para(tf, "• Piloto A/B calibra as premissas (efeito, resposta, margem) antes de qualquer escala", 14, TEXT_MAIN)
    _kpi(s, 7.0, 1.6, "22", "tabelas com contrato validado")
    _kpi(s, 10.0, 1.6, "190", "testes automatizados verdes")
    _kpi(s, 7.0, 3.3, "R$ 0", "custo do agente offline por recomendação", accent=LIME)
    _kpi(s, 10.0, 3.3, "9", "workstreams entregues com revisão independente")
    _footer(s, "Cenário de custo LLM: estimativa prospectiva rotulada como premissa (600–900 tokens/rec, câmbio 5,50).")

    # ---------------- 8 · FECHAMENTO ----------------
    s = _slide(prs)
    _bg(s, AZUL_DEEP)
    _stripe(s, y=0.55, color=MAGENTA)
    tf = _box(s, 0.9, 1.8, 11.5, 1.4)
    _para(tf, "Da milha ao momento de cuidado", 40, WHITE, bold=True)
    tf = _box(s, 0.9, 3.3, 11.5, 2.6)
    _para(tf, "1. Validar em piloto A/B com grupo de controle (humano no circuito, opt-out, monitoramento de viés)", 18, RGBColor(0x9A, 0xE0, 0xE8))
    _para(tf, "2. Calibrar efeito/resposta/margem com uplift real — substituir premissas por medidas", 18, RGBColor(0x9A, 0xE0, 0xE8))
    _para(tf, "3. Conectar os dados reais no Snowflake e re-auditar (fairness a cada re-treino)", 18, RGBColor(0x9A, 0xE0, 0xE8))
    _para(tf, "Demo ao vivo: python tasks.py app · Painel em http://localhost:8501", 16, LIME)
    _footer(s, "Projeto Aura 2.0 · Entrega com evidências executadas — 12 revisões independentes, 190 testes, zero número inventado.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"deck salvo: {OUT} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)" if False else f"deck salvo: {OUT}")


if __name__ == "__main__":
    build()
